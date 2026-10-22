import io
import pytest
from docx import Document

from app.models.schemas import (
    FullAnalysisReport,
    TenderSummary,
    MedComplianceAudit,
    TrafficLightVerdict,
    RiskItem,
    RiskSeverity,
    RiskCategory,
    EquipmentItem,
    ClarificationRequest,
)
from app.utils.export import ReportExporter


@pytest.fixture
def sample_report():
    return FullAnalysisReport(
        verdict=TrafficLightVerdict.RED,
        risk_score=85,
        verdict_summary="Выявлены кабальные условия штрафов и заточка под бренд.",
        summary=TenderSummary(
            purchase_subject="Поставка ультразвуковой диагностической системы",
            law_type="44-ФЗ",
            notice_number="0142200001326018834",
            customer_name='ГБУЗ "Самарская областная клиническая больница"',
            customer_signer="Главный врач Иванов И.И.",
            customer_inn="6315001234",
            customer_kpp="631501001",
            customer_address="г. Самара, ул. Ленина, д. 1",
            nmck_rubles=15000000.0,
            delivery_deadline="30 календарных дней",
            advance_percent=0.0,
            contract_guarantee_percent=5.0,
            etp_name="ЕЭТП",
        ),
        med_compliance=MedComplianceAudit(
            ru_required=True,
            ru_notes="Требуется действующее РУ",
            national_regime_applies=True,
            national_regime_details="ПП 878 радиоэлектроника",
            ktru_used=True,
            ktru_code="26.60.12.119-00000001",
            warranty_period_months=24,
            installation_and_training=True,
        ),
        risks=[
            RiskItem(
                title="Штрафы выше нормы ПП 1042",
                description="Штраф 10% вместо 5%",
                severity=RiskSeverity.HIGH,
                category=RiskCategory.LEGAL,
                quote="Штраф 10%",
                law_reference="ПП РФ № 1042",
                recommendation="Требовать снижения",
            )
        ],
        equipment=[
            EquipmentItem(
                name="Ультразвуковой аппарат экспертного класса",
                quantity="1 шт.",
                key_specs=["Датчики 3 шт.", "Монитор 21.5 дюймов"],
                potential_bias_warning="Заточка под Mindray Resona 7",
                detected_brand_bias="Эксклюзивная технология зонного сканирования",
                suggested_neutral_specs="Заменить на эквивалентную формулировку",
            )
        ],
        clarification_requests=[
            ClarificationRequest(
                subject="О снижении размера штрафа до норматива ПП РФ № 1042",
                justification="ч. 4 ст. 34 44-ФЗ и ПП РФ № 1042",
                draft_text="Просим привести размер штрафа в соответствие с законодательством...",
                recipient_header="Главному врачу ГБУЗ СОКБ",
            )
        ],
    )


def test_export_docx_report(sample_report):
    docx_bytes = ReportExporter.to_docx_report(sample_report)
    assert len(docx_bytes) > 0
    # Проверяем, что сгенерированный файл корректно открывается python-docx
    doc = Document(io.BytesIO(docx_bytes))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "ЭКСПЕРТНОЕ ЗАКЛЮЧЕНИЕ" in full_text
    assert "Самарская областная клиническая больница" in " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_export_docx_clarification_request(sample_report):
    cr = sample_report.clarification_requests[0]
    docx_bytes = ReportExporter.to_docx_clarification_request(cr, sample_report.summary, 1)
    assert len(docx_bytes) > 0
    doc = Document(io.BytesIO(docx_bytes))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "ЗАПРОС О ДАЧЕ РАЗЪЯСНЕНИЙ" in full_text
    assert "ПП РФ № 1042" in full_text


def test_export_docx_all_clarifications(sample_report):
    docx_bytes = ReportExporter.to_docx_all_clarifications(
        sample_report.clarification_requests, sample_report.summary
    )
    assert len(docx_bytes) > 0
    doc = Document(io.BytesIO(docx_bytes))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "ЗАПРОС" in full_text


def test_export_docx_fas_complaint(sample_report):
    from app.services.fas_generator import FASComplaintGenerator
    complaint = FASComplaintGenerator.generate(sample_report, applicant_name="ООО МедТех")
    assert complaint.fas_authority_name is not None
    assert "Самарск" in complaint.fas_authority_name
    assert len(complaint.episodes) >= 1
    assert len(complaint.petitum) >= 3

    docx_bytes = ReportExporter.export_fas_complaint_docx(complaint)
    assert len(docx_bytes) > 0
    doc = Document(io.BytesIO(docx_bytes))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "ЖАЛОБА" in full_text
    assert "ПРОШУ" in full_text
    assert "ООО МедТех" in " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
